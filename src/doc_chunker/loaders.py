from dataclasses import dataclass
from pathlib import Path
from typing import List, Optional

from .errors import NoTextError, UnsupportedFileError


@dataclass
class Section:
    """A piece of text with its page or slide number, if there is one."""

    text: str
    page: Optional[int] = None


TEXT_SUFFIXES = {".txt", ".md", ".markdown"}
SUPPORTED = TEXT_SUFFIXES | {".pdf", ".docx", ".pptx"}


def load(path: Path) -> List[Section]:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        raise UnsupportedFileError(
            f"{path.name}: file type {suffix or '(none)'} is not supported. "
            f"Supported: {', '.join(sorted(SUPPORTED))}."
        )
    if not path.is_file():
        raise UnsupportedFileError(f"{path.name}: file not found.")
    try:
        if suffix == ".pdf":
            sections = _load_pdf(path)
        elif suffix == ".docx":
            sections = _load_docx(path)
        elif suffix == ".pptx":
            sections = _load_pptx(path)
        else:
            sections = [Section(_read_text(path))]
    except (UnsupportedFileError, NoTextError):
        raise
    except Exception as exc:
        raise UnsupportedFileError(f"{path.name}: file cannot be read ({exc}).") from exc

    sections = [s for s in sections if s.text.strip()]
    if not sections:
        hint = ""
        if suffix == ".pdf":
            hint = " For scans, use doc-extract, which reads the text with a vision model."
        raise NoTextError(f"{path.name}: no extractable text.{hint}")
    return sections


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("latin-1")


def _load_pdf(path: Path) -> List[Section]:
    from pdfminer.high_level import extract_pages
    from pdfminer.layout import LAParams
    from pdfminer.pdfdocument import PDFPasswordIncorrect

    try:
        pages = list(extract_pages(str(path), laparams=LAParams(boxes_flow=None)))
    except PDFPasswordIncorrect:
        raise UnsupportedFileError(f"{path.name}: PDF is password protected.") from None
    return [Section(_pdf_page_text(page), number) for number, page in enumerate(pages, 1)]


def _pdf_page_text(page) -> str:
    """Text of one page. Text blocks are put in reading order from their positions."""
    from pdfminer.layout import LTTextContainer

    blocks = []
    for element in page:
        if isinstance(element, LTTextContainer):
            text = "\n".join(line.rstrip() for line in element.get_text().splitlines()).strip()
            if text:
                blocks.append((element.x0, element.y0, element.x1, element.y1, text))
    unit = max(page.width, page.height) * 0.01
    return "\n\n".join(b[4] for b in _reading_order(blocks, unit))


def _reading_order(blocks: list, unit: float) -> list:
    """Recursive XY cut: split at the widest empty gap, either vertical (between columns) or
    horizontal (below a heading, between sections). Blocks that cannot be split are read top to
    bottom, left to right. Using the widest gap first keeps columns together, because the gap
    between columns is wider than the gap between paragraphs.

    A block is (x0, y0, x1, y1, text) with y growing upwards. `unit` is the smallest gap
    that counts as a cut.
    """
    if len(blocks) < 2:
        return list(blocks)
    columns = _widest_cut(blocks, unit, lo=0, hi=2, descending=False)
    rows = _widest_cut(blocks, unit, lo=1, hi=3, descending=True)
    cuts = [c for c in (columns, rows) if c is not None]
    if not cuts:
        return sorted(blocks, key=lambda b: (-b[3], b[0]))
    _, first, second = max(cuts, key=lambda c: c[0])
    return _reading_order(first, unit) + _reading_order(second, unit)


def _widest_cut(blocks: list, unit: float, lo: int, hi: int, descending: bool):
    """Find the widest empty gap along one axis. Returns (gap, before, after) or None."""
    ordered = sorted(blocks, key=lambda b: b[hi] if descending else b[lo], reverse=descending)
    best = None
    edge = None
    for i, b in enumerate(ordered):
        if edge is not None:
            gap = edge - b[hi] if descending else b[lo] - edge
            if gap >= unit and (best is None or gap > best[0]):
                best = (gap, i)
        far = b[lo] if descending else b[hi]
        edge = far if edge is None else (min(edge, far) if descending else max(edge, far))
    if best is None:
        return None
    gap, i = best
    return gap, ordered[:i], ordered[i:]


def _load_docx(path: Path) -> List[Section]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))
    blocks: List[str] = []
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            text = para.text.strip()
            if not text:
                continue
            style = para.style.name if para.style is not None else ""
            if style.startswith("Heading ") and style[8:].isdigit():
                text = "#" * min(int(style[8:]), 6) + " " + text
            elif style == "Title":
                text = "# " + text
            blocks.append(text)
        elif tag == "tbl":
            blocks.append(_table_text(Table(child, document)))
    return [Section("\n\n".join(b for b in blocks if b))]


def _table_text(table) -> str:
    rows = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            value = cell.text.strip().replace("\n", " ")
            if not cells or cells[-1] != value:
                cells.append(value)
        rows.append(" | ".join(cells))
    return "\n".join(rows)


def _load_pptx(path: Path) -> List[Section]:
    from pptx import Presentation

    presentation = Presentation(str(path))
    sections = []
    for number, slide in enumerate(presentation.slides, 1):
        lines: List[str] = []
        _shape_lines(slide.shapes, lines)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                lines.append("Notes: " + notes.text.strip())
        sections.append(Section("\n".join(lines), number))
    return sections


def _shape_lines(shapes, lines: List[str]) -> None:
    for shape in shapes:
        if shape.shape_type == 6:  # group
            _shape_lines(shape.shapes, lines)
        elif getattr(shape, "has_table", False) and shape.has_table:
            lines.append(_table_text(shape.table))
        elif shape.has_text_frame and shape.text_frame.text.strip():
            lines.append(shape.text_frame.text.strip())
