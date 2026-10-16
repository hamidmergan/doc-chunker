import json

import pytest

from doc_chunker import ConfigError, NoTextError, UnsupportedFileError, chunk, chunk_many
from doc_chunker.cli import main
from doc_chunker.splitter import split_text


def test_split_respects_size_and_overlap():
    text = " ".join(f"word{i}" for i in range(400))
    parts = split_text(text, 200, 40)
    assert len(parts) > 1
    assert all(len(p) <= 200 for p in parts)
    assert parts[0].split()[-1] in parts[1]


def test_split_hard_cut_without_separators():
    parts = split_text("x" * 450, 100, 0)
    assert [len(p) for p in parts] == [100, 100, 100, 100, 50]


def test_txt_short_is_one_chunk(tmp_path):
    f = tmp_path / "a.txt"
    f.write_text("Hello world", encoding="utf-8")
    result = chunk(f)
    assert len(result) == 1
    assert result[0].text == "Hello world"
    assert result[0].source == "a.txt"
    assert result[0].page is None


def test_markdown_heading_starts_new_chunk(tmp_path):
    f = tmp_path / "a.md"
    f.write_text("# One\n\nText one.\n\n# Two\n\nText two.\n", encoding="utf-8")
    result = chunk(f)
    assert [c.text.split("\n")[0] for c in result] == ["# One", "# Two"]
    assert [c.index for c in result] == [0, 1]


def test_latin1_text(tmp_path):
    f = tmp_path / "a.txt"
    f.write_bytes("Grüße".encode("latin-1"))
    assert chunk(f)[0].text == "Grüße"


def test_empty_file_raises(tmp_path):
    f = tmp_path / "a.txt"
    f.write_text("   \n", encoding="utf-8")
    with pytest.raises(NoTextError):
        chunk(f)


def test_unsupported_suffix(tmp_path):
    f = tmp_path / "a.xyz"
    f.write_text("x")
    with pytest.raises(UnsupportedFileError):
        chunk(f)


def test_missing_file(tmp_path):
    with pytest.raises(UnsupportedFileError):
        chunk(tmp_path / "missing.txt")


def test_invalid_settings(tmp_path):
    f = tmp_path / "a.txt"
    f.write_text("x" * 100)
    with pytest.raises(ConfigError):
        chunk(f, chunk_size=10)
    with pytest.raises(ConfigError):
        chunk(f, chunk_size=100, overlap=100)


def test_env_settings(tmp_path, monkeypatch):
    f = tmp_path / "a.txt"
    f.write_text("word " * 100)
    monkeypatch.setenv("DOC_CHUNKER_CHUNK_SIZE", "100")
    monkeypatch.setenv("DOC_CHUNKER_OVERLAP", "0")
    assert all(len(c.text) <= 100 for c in chunk(f))
    monkeypatch.setenv("DOC_CHUNKER_CHUNK_SIZE", "abc")
    with pytest.raises(ConfigError):
        chunk(f)


def test_docx(tmp_path):
    import docx

    d = docx.Document()
    d.add_heading("Title", level=1)
    d.add_paragraph("First paragraph.")
    d.add_heading("Part", level=2)
    d.add_paragraph("Second paragraph.")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "A"
    table.rows[0].cells[1].text = "B"
    f = tmp_path / "a.docx"
    d.save(f)
    result = chunk(f)
    assert result[0].text.startswith("# Title")
    assert result[1].text.startswith("## Part")
    assert "A | B" in result[1].text


def test_pptx_one_chunk_per_slide(tmp_path):
    from pptx import Presentation

    p = Presentation()
    for title in ("Slide one", "Slide two"):
        s = p.slides.add_slide(p.slide_layouts[5])
        s.shapes.title.text = title
    f = tmp_path / "a.pptx"
    p.save(f)
    result = chunk(f)
    assert [c.page for c in result] == [1, 2]
    assert result[1].text == "Slide two"


def test_pdf_pages_and_scan_message(tmp_path):
    blank = tmp_path / "blank.pdf"
    blank.write_bytes(_pdf_with_pages([""]))
    with pytest.raises(NoTextError, match="doc-extract"):
        chunk(blank)

    f = tmp_path / "a.pdf"
    f.write_bytes(_pdf_with_pages(["Page one", "Page two"]))
    result = chunk(f)
    assert [c.page for c in result] == [1, 2]
    assert "one" in result[0].text


def _pdf_with_pages(texts):
    objs = []
    n = len(texts)
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{3 + i} 0 R" for i in range(n))
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {n} >>".encode())
    font_id = 3 + 2 * n
    for i in range(n):
        objs.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 300] "
            f"/Contents {4 + n + i - 1 + 1 - 1} 0 R /Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode()
        )
    for t in texts:
        stream = f"BT /F1 18 Tf 20 150 Td ({t}) Tj ET".encode()
        objs.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = b"%PDF-1.4\n"
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n".encode() + b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


def test_chunk_many_isolates_errors(tmp_path):
    good = tmp_path / "a.txt"
    good.write_text("ok")
    bad = tmp_path / "b.xyz"
    bad.write_text("x")
    results = chunk_many([good, bad])
    assert results[str(good)].ok
    assert not results[str(bad)].ok


def test_cli_stdout_and_exit_codes(tmp_path, capsys):
    f = tmp_path / "a.txt"
    f.write_text("Hello")
    assert main([str(f)]) == 0
    row = json.loads(capsys.readouterr().out.strip())
    assert row == {"source": "a.txt", "index": 0, "page": None, "text": "Hello"}

    assert main([str(tmp_path / "x.xyz")]) == 1
    assert main([str(f), "--chunk-size", "10"]) == 2


def test_cli_output_file_and_folder(tmp_path):
    a = tmp_path / "a.txt"
    b = tmp_path / "b.txt"
    a.write_text("one")
    b.write_text("two")
    one = tmp_path / "out.jsonl"
    assert main([str(a), "-o", str(one)]) == 0
    assert json.loads(one.read_text())["text"] == "one"
    folder = tmp_path / "folder"
    assert main([str(a), str(b), "-o", str(folder)]) == 0
    assert (folder / "a.txt.jsonl").exists()
    assert (folder / "b.txt.jsonl").exists()


def _order(blocks):
    from doc_chunker.loaders import _reading_order

    return [b[4] for b in _reading_order(blocks, unit=6)]


def _block(x0, y0, x1, y1, text):
    return (x0, y0, x1, y1, text)


def test_two_columns_are_read_column_by_column():
    blocks = [
        _block(50, 700, 250, 720, "L1"), _block(300, 700, 500, 720, "R1"),
        _block(50, 650, 250, 670, "L2"), _block(300, 650, 500, 670, "R2"),
        _block(50, 600, 250, 620, "L3"), _block(300, 600, 500, 620, "R3"),
    ]
    assert _order(blocks) == ["L1", "L2", "L3", "R1", "R2", "R3"]


def test_heading_across_columns_comes_first_and_footer_last():
    blocks = [
        _block(50, 760, 500, 780, "Title"),
        _block(50, 700, 250, 720, "L1"), _block(300, 700, 500, 720, "R1"),
        _block(50, 650, 250, 670, "L2"), _block(300, 650, 500, 670, "R2"),
        _block(50, 40, 500, 60, "Footer"),
    ]
    assert _order(blocks) == ["Title", "L1", "L2", "R1", "R2", "Footer"]


def test_single_column_is_top_to_bottom():
    blocks = [_block(50, y, 500, y + 20, f"P{i}") for i, y in enumerate([600, 700, 500])]
    assert _order(blocks) == ["P1", "P0", "P2"]


def test_three_columns():
    blocks = [_block(x, y, x + 100, y + 20, f"{c}{n}") for c, x in zip("ABC", (50, 200, 350)) for n, y in ((1, 700), (2, 600))]
    blocks += [_block(x, 650, x + 100, 670, f"{c}x") for c, x in zip("ABC", (50, 200, 350))]
    assert _order(blocks) == ["A1", "Ax", "A2", "B1", "Bx", "B2", "C1", "Cx", "C2"]
